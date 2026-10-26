"""Draw rooms onto 16:9 slides, each room in its real layout and orientation.

draw_room(prs, room, title, subtitle, seat_text, style) — seat_text: {desk: [(text, role)]}, style: {desk: 'used'|'warn'}
labels: {'podium': text on the podium, 'blocked': text on blocked cells (default: the room config's label)}
"""
import math
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

W, H, M = 960, 540, 14
FONT = 'Microsoft JhengHei'
INK, SUB, LINE, EMPTY = RGBColor(0x1F, 0x1F, 0x1F), RGBColor(0x55, 0x55, 0x55), RGBColor(0x59, 0x59, 0x59), RGBColor(0xBF, 0xBF, 0xBF)
WHITE, PILLAR = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xD0, 0xD0, 0xD0)
ACCENT, USED, WARN = RGBColor(0xF0, 0xA0, 0x30), RGBColor(0x4A, 0x7B, 0xC8), RGBColor(0xD9, 0x4A, 0x38)
TITLE_W, TITLE_H, BAND = 320, 54, 62

def text(tf, runs, align=PP_ALIGN.CENTER):
    tf.clear(); tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, (s, size, bold, color) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = s
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = color

def box(slide, x, y, w, h, fill=WHITE, line=LINE, lw=1.25, rot=0.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    sh.shadow.inherit = False
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.rotation = rot
    return sh

def ems(s):
    """Rough text width in em: CJK characters are full width, digits 0.55, Latin letters and the rest 0.6."""
    return sum(1.0 if ord(c) >= 0x2E80 else 0.55 if c.isdigit() else 0.6 for c in s)

def hit(a, b):
    return a[0] < b[0] + b[2] and b[0] < a[0] + a[2] and a[1] < b[1] + b[3] and b[1] < a[1] + a[3]

class Fit:
    """Scale the room (all desks + podium) into a slide area; optional horizontal stretch."""
    def __init__(self, room, top):
        self.sx = room.stretch_x
        pts = [p for d in room.desks for p in self.corners(d)] + self.corners(dict(room.podium, rot=0))
        x0, x1 = min(p[0] for p in pts), max(p[0] for p in pts)
        y0, y1 = min(p[1] for p in pts), max(p[1] for p in pts)
        self.k = min((W - 2 * M) / (x1 - x0), (H - M - top) / (y1 - y0))
        self.ox = (W - self.k * (x1 - x0)) / 2 - self.k * x0
        self.oy = top + (H - M - top - self.k * (y1 - y0)) / 2 - self.k * y0

    def corners(self, d):
        a = math.radians(d['rot']); c, s = math.cos(a), math.sin(a)
        return [(d['cx'] * self.sx + u * c - v * s, d['cy'] + u * s + v * c)
                for u in (-d['w'] / 2, d['w'] / 2) for v in (-d['h'] / 2, d['h'] / 2)]

    def centre(self, d):
        return self.ox + self.k * d['cx'] * self.sx, self.oy + self.k * d['cy']

    def bounds(self, d):
        xs, ys = zip(*[(self.ox + self.k * x, self.oy + self.k * y) for x, y in self.corners(d)])
        return min(xs), min(ys), max(xs) - min(xs), max(ys) - min(ys)

def layout(room):
    """Fit the room to the whole slide and put the title in free space at the top; if there is none, keep a band for it."""
    fit = Fit(room, M)
    taken = [fit.bounds(d) for d in room.desks] + [fit.bounds(dict(room.podium, rot=0))]
    for x in (M, (W - TITLE_W) / 2, W - M - TITLE_W):
        spot = (x, 6, TITLE_W, TITLE_H)
        if not any(hit(spot, t) for t in taken):
            return fit, spot
    fit = Fit(room, BAND)                             # no room at the top: keep a band, title left of the podium
    pod_left = fit.bounds(dict(room.podium, rot=0))[0]
    return fit, (M, 6, max(TITLE_W, pod_left - M - 10) if pod_left > M + 160 else W - 2 * M, TITLE_H)

def label_size(fit, room, seat_text, cap=24):
    """Largest name size where no two seat labels collide (and, on straight desks, each fits its desk)."""
    def boxes(npt):
        ipt = max(6.5, round(npt * 0.6 * 2) / 2)
        out = []
        for i, lines in seat_text.items():
            cx, cy = fit.centre(room.desks[i])
            w = max(ems(s) * (npt if role == 'main' else ipt) for s, role in lines) + 1
            h = sum(npt * 1.2 if role == 'main' else ipt * 1.15 for _, role in lines)
            out.append((i, (cx - w / 2, cy - h / 2, w, h)))
        return out, ipt
    pod = fit.bounds(dict(room.podium, rot=0))
    for npt in range(cap, 6, -1):
        bx, ipt = boxes(npt)
        rects = [b for _, b in bx]
        clash = any(hit(a, b) for n, a in enumerate(rects) for b in rects[n + 1:]) or any(hit(a, pod) for a in rects)
        too_big = any((b[2] > 0.92 * fit.k * room.desks[i]['w'] or b[3] > fit.k * room.desks[i]['h'])
                      if room.desks[i]['rot'] % 90 == 0 else b[2] > 1.5 * fit.k * room.desks[i]['w']   # slanted desks: stay near the desk
                      for i, b in bx)
        too_big = too_big or any(b[0] < 2 or b[1] < 2 or b[0] + b[2] > W - 2 or b[1] + b[3] > H - 16 for _, b in bx)
        if not clash and not too_big:
            return npt, ipt, dict(bx)
    bx, ipt = boxes(7)
    return 7, ipt, dict(bx)

def draw_room(prs, room, title, subtitle, seat_text=None, style=None, footer=None, cap=24, labels=None):
    seat_text, style, labels = seat_text or {}, style or {}, labels or {}
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fit, spot = layout(room)
    x, y, w, h = fit.bounds(dict(room.podium, rot=0))
    text(box(slide, x, y, w, h, fill=ACCENT, line=None).text_frame, [(labels.get('podium', 'Podium'), 20, True, INK)])
    t = slide.shapes.add_textbox(*map(Pt, spot))
    runs = [(title, 28, True, INK)] + ([(subtitle, 13, False, SUB)] if subtitle else [])
    text(t.text_frame, runs, PP_ALIGN.LEFT if spot[0] == M else PP_ALIGN.CENTER)
    for i, d in enumerate(room.desks):
        cx, cy = fit.centre(d)
        w, h = fit.k * d['w'], fit.k * d['h']
        if 'blocked' in d:
            sh = box(slide, cx - w / 2, cy - h / 2, w, h, fill=PILLAR, line=None, rot=d['rot'])
            text(sh.text_frame, [(labels.get('blocked', d['blocked']), 14, True, SUB)])
            continue
        kind = style.get(i)
        if kind:
            box(slide, cx - w / 2, cy - h / 2, w, h, fill=USED, line=WARN if kind == 'warn' else None, lw=2.5, rot=d['rot'])
        elif i in seat_text:
            box(slide, cx - w / 2, cy - h / 2, w, h, rot=d['rot'])
        else:
            box(slide, cx - w / 2, cy - h / 2, w, h, line=EMPTY, lw=1, rot=d['rot'])
    if seat_text:
        npt, ipt, rects = label_size(fit, room, seat_text, cap)
        colour = {'main': INK, 'sub': SUB, 'code': SUB}
        for i, lines in seat_text.items():
            on_blue = style.get(i) is not None
            tb = slide.shapes.add_textbox(*map(Pt, rects[i]))
            text(tb.text_frame, [(s, npt if role == 'main' else ipt, role == 'main',
                                  WHITE if on_blue else colour[role]) for s, role in lines])
    if footer:
        t = slide.shapes.add_textbox(Pt(M), Pt(H - 14), Pt(W - 2 * M), Pt(12))
        text(t.text_frame, [(footer, 9, False, SUB)], PP_ALIGN.LEFT)
    return slide
