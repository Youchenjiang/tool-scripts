"""Arrange exam seats from an exam config.

    python src/seating.py exams/<exam>.json

Reads the roster and the rooms named in the config (rooms/<room>.json), then
  stage 1 (which seats are used) -> output/<exam>_<course>_seat-layout-preview.pptx: used seats in blue, every seat marked
                                   with its code (group letter + row + seat, e.g. B1-2) for the config's "fixed" section
  stage 2 (who sits where)       -> output/<exam>_<course>_seating-chart.pptx: the chart shown to students (room name only);
                                   the rule check is printed to the terminal
"""
import json, os, sys
import openpyxl
from pptx import Presentation
from pptx.util import Pt
import rooms, solver, render, assign

def load_roster(path, name_col='Name', id_col='ID'):
    ws = openpyxl.load_workbook(path).active
    rows = list(ws.iter_rows(values_only=True))
    head = next(n for n, r in enumerate(rows) if name_col in r and id_col in r)
    ni, ii = rows[head].index(name_col), rows[head].index(id_col)
    people = [dict(name=str(r[ni]).strip(), sid=str(r[ii]).strip()) for r in rows[head + 1:] if r[ii]]
    return sorted(people, key=lambda p: p['sid'])

def main(config_path):
    cfg = json.load(open(config_path, encoding='utf-8'))
    base = os.path.dirname(os.path.abspath(config_path))
    people = load_roster(os.path.join(base, cfg['roster']), *cfg.get('roster_columns', ['Name', 'ID']))
    room_list = [rooms.load(name) for name in cfg['rooms']]
    by_name = {r.name: ri for ri, r in enumerate(room_list)}

    fixed = {}
    for name, spec in cfg.get('fixed', {}).items():
        room = room_list[by_name[name]]
        for code in spec.get('sit', []):
            fixed[(by_name[name], room.index[code])] = True
        for code in spec.get('empty', []):
            fixed[(by_name[name], room.index[code])] = False
    counts = {by_name[n]: c for n, c in cfg.get('counts', {}).items()}

    chosen = solver.stage1(room_list, len(people), fixed, counts)
    side, diag, front = (solver.pairs(room_list, chosen, k) for k in ('side', 'diag', 'front'))
    print(f'{len(people)} students | ' + ' | '.join(f'{r.name}: {sum(1 for ri, _ in chosen if ri == n)}' for n, r in enumerate(room_list))
          + f' | side {side} diag {diag} front {front}')

    prs = new_deck()
    footer = (f'{cfg["course"]}   blue = occupied   red border = diagonal neighbour   '
              f'side-by-side {side}, diagonal {diag}, front/back {front} pairs   code = block, row - seat')
    for ri, room in enumerate(room_list):
        used = {i for rj, i in chosen if rj == ri}
        style = {i: 'warn' if room.nbr[i]['diag'] & used else 'used' for i in used}
        codes = {i: [(room.label[i], 'code')] for i in room.seats}
        render.draw_room(prs, room, f'Room {room.name}', f'Seat layout preview   {len(used)} students', codes, style, footer, cap=16)
    save(prs, base, cfg, 'seat-layout-preview')

    if not cfg.get('rules'):
        return
    placed = assign.assign(room_list, chosen, people, cfg['rules'], cfg.get('seed', 1))
    results = assign.check(room_list, chosen, people, cfg['rules'], placed)
    for rule, ok, detail in results:                  # checks stay in the terminal: the chart is shown to students
        print(f'  [{"OK" if ok else "NG"}] {rule}: {detail}')
    prs = new_deck()
    for ri, room in enumerate(room_list):
        names = {i: [(people[p]['name'], 'main'), (people[p]['sid'], 'sub')] for (rj, i), p in placed.items() if rj == ri}
        render.draw_room(prs, room, f'Room {room.name}', '', names)
    save(prs, base, cfg, 'seating-chart')

def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(render.W), Pt(render.H)
    return prs

def save(prs, base, cfg, kind):
    out = os.path.join(base, cfg['output'], f'{cfg["name"]}_{cfg["course"]}_{kind}.pptx')
    os.makedirs(os.path.dirname(out), exist_ok=True)
    prs.save(out)
    print('saved', os.path.normpath(out))

if __name__ == '__main__':
    main(sys.argv[1])
