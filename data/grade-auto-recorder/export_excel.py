#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
成績單 Excel 導出模組 (CSV -> 美化後的 XLSX)
提供專業美觀的排版、自動欄寬、標題色彩與格線。
"""

import csv
import sys
from pathlib import Path

# 確保在 Windows 控制台正確輸出
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_CSV = BASE_DIR / "grades.csv"
DEFAULT_XLSX = BASE_DIR / "grades.xlsx"


def convert_csv_to_excel(csv_path=DEFAULT_CSV, xlsx_path=DEFAULT_XLSX):
    csv_path = Path(csv_path)
    xlsx_path = Path(xlsx_path)

    if not csv_path.exists():
        raise FileNotFoundError(f"找不到 CSV 檔案: {csv_path}")

    # 讀取 CSV 資料
    rows = []
    with open(csv_path, mode="r", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        rows = [row for row in reader if any(cell.strip() for cell in row)]

    if not rows:
        print("[WARN] CSV 檔案為空，未產生 Excel。")
        return None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "學生成績表"

    # 樣式定義
    header_font = Font(name="微軟正黑體", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")  # 沉穩深藍
    data_font = Font(name="微軟正黑體", size=10)
    
    # 斑馬紋交替色
    zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    # 寫入資料與套用樣式
    for r_idx, row in enumerate(rows, start=1):
        is_header = (r_idx == 1)
        is_zebra = (r_idx % 2 == 0) and not is_header

        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border

            if is_header:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
            else:
                cell.font = data_font
                cell.fill = zebra_fill if is_zebra else white_fill
                # 帳號、成績等欄位置中，其他靠左或置中
                cell.alignment = center_align

    # 凍結首行
    ws.freeze_panes = "A2"

    # 自動調整欄寬
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            # 計算中文字元寬度（中文字算 2，英數算 1）
            width = sum(2 if ord(c) > 127 else 1 for c in val_str)
            if width > max_len:
                max_len = width
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(xlsx_path)
    print(f"[SUCCESS] 已自動導出 Excel: {xlsx_path} (共 {len(rows)-1} 筆學生資料)")
    return xlsx_path


if __name__ == "__main__":
    if len(sys.argv) > 1:
        convert_csv_to_excel(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else Path(sys.argv[1]).with_suffix(".xlsx"))
    else:
        # 自動轉換 data/ 或根目錄下所有的 grade_*.csv 與 grades.csv
        search_dirs = [BASE_DIR / "data", BASE_DIR]
        processed = set()
        for d in search_dirs:
            if d.exists():
                for p in d.glob("grade_*.csv"):
                    if p not in processed:
                        convert_csv_to_excel(p, p.with_suffix(".xlsx"))
                        processed.add(p)
                for p in d.glob("grades.csv"):
                    if p not in processed:
                        convert_csv_to_excel(p, p.with_suffix(".xlsx"))
                        processed.add(p)
