#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
學生成績管理系統 - 檔案匯出模組 (Exporter)
支援 CSV 儲存與專業美化版 Excel (openpyxl) 導出（自動欄寬、凍結首行、斑馬紋）。
"""

import csv
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import Roster


class GradeExporter:
    """成績檔案匯出器"""

    @staticmethod
    def save_to_csv(
        roster: Roster,
        grades_map: Dict[str, str],
        output_dir: Path,
        custom_name: Optional[str] = None,
        column_name: str = "成績"
    ) -> Path:
        """將成績寫入 CSV 檔案"""
        output_dir = Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)

        today_str = datetime.now().strftime("%Y%m%d")
        file_prefix = custom_name.strip() if custom_name else f"grade_{today_str}"
        csv_path = output_dir / f"{file_prefix}.csv"

        rows = []
        for s in roster.students:
            score = grades_map.get(s.sid, "")
            rows.append({
                "帳號": s.sid,
                "姓名": s.name,
                column_name: score,
                "備註": ""
            })

        fieldnames = ["帳號", "姓名", column_name, "備註"]
        with open(csv_path, mode="w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)

        return csv_path

    @staticmethod
    def convert_csv_to_excel(csv_path: Path, xlsx_path: Optional[Path] = None) -> Path:
        """將 CSV 轉為具備深藍表頭、最適欄寬、斑馬紋交替色的 Excel 檔案"""
        csv_path = Path(csv_path)
        xlsx_path = Path(xlsx_path) if xlsx_path else csv_path.with_suffix(".xlsx")

        if not csv_path.exists():
            raise FileNotFoundError(f"找不到 CSV 檔案: {csv_path}")

        rows = []
        with open(csv_path, mode="r", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            rows = [row for row in reader if any(cell.strip() for cell in row)]

        if not rows:
            raise ValueError("CSV 檔案為空，無法產生 Excel")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "學生成績表"

        header_font = Font(name="微軟正黑體", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        data_font = Font(name="微軟正黑體", size=10)

        zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        center_align = Alignment(horizontal="center", vertical="center")

        for r_idx, row in enumerate(rows, start=1):
            is_header = (r_idx == 1)
            is_zebra = (r_idx % 2 == 0) and not is_header

            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.border = thin_border

                if is_header:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center_align
                else:
                    cell.font = data_font
                    cell.fill = zebra_fill if is_zebra else white_fill
                    cell.alignment = center_align

        ws.freeze_panes = "A2"

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                width = sum(2 if ord(c) > 127 else 1 for c in val_str)
                if width > max_len:
                    max_len = width
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(xlsx_path)
        return xlsx_path

    @classmethod
    def export_full(
        cls,
        roster: Roster,
        grades_map: Dict[str, str],
        output_dir: Path,
        custom_name: Optional[str] = None,
        column_name: str = "成績"
    ) -> Tuple[Path, Path]:
        """同時匯出 CSV 與美化版 Excel"""
        csv_path = cls.save_to_csv(roster, grades_map, output_dir, custom_name, column_name)
        xlsx_path = cls.convert_csv_to_excel(csv_path)
        return csv_path, xlsx_path
